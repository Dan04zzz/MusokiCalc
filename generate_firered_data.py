import openpyxl
import re
import json

SPECIES_CORRECTIONS = {
    "Skamory": "Skarmory",
    "Snubbul": "Snubbull",
    "Farfetchd": "Farfetch'd",
    "Sirfetchd": "Sirfetch'd",
    "Mr-Mime": "Mr. Mime",
    "Mime-Jr": "Mime Jr.",
    "Mr-Rime": "Mr. Rime",
}

MOVE_CORRECTIONS = {
    "ExtremeSpeed": "Extreme Speed",
    "Poisonpowder": "Poison Powder",
    "Sonicboom": "Sonic Boom",
    "ThunderPunch": "Thunder Punch",
    "Thundershock": "Thunder Shock",
}

ITEM_CORRECTIONS = {
    "Never-Melt Ice": "NeverMelt Ice",
}

def clean_species(name):
    if not name:
        return ""
    name = str(name).strip()
    # Remove owner suffix like (A) or (G)
    name = re.sub(r"\s*[\(\[].*?[\)\]]", "", name)
    name = name.strip()
    return SPECIES_CORRECTIONS.get(name, name)

def clean_move(name):
    if not name:
        return ""
    name = str(name).strip()
    if name == "-" or name.lower() == "none" or name.lower() == "null":
        return ""
    return MOVE_CORRECTIONS.get(name, name)

def clean_item(name):
    if not name:
        return "None"
    name = str(name).strip()
    if name.lower() == "none" or name.lower() == "null" or name == "-":
        return "None"
    return ITEM_CORRECTIONS.get(name, name)

def clean_nature(name):
    if not name:
        return "Serious"
    name = str(name).strip()
    return name

def clean_ability(name):
    if not name:
        return "None"
    return str(name).strip()

def main():
    wb = openpyxl.load_workbook("Documento Lotte Pokemon Rosso Fuoco 2.0.xlsx", read_only=True)
    
    # We map substrings of target sheet names to handle encoding issues robustly
    sheet_mapping = {
        "gym leader": None,
        "rival": None,
        "team rocket": None,
        "league": None
    }
    
    for actual_name in wb.sheetnames:
        for key in sheet_mapping.keys():
            if key in actual_name.lower():
                sheet_mapping[key] = actual_name
                break
                
    sheets = [
        ("gym leader", "Gym Leader"),
        ("rival", "Rival"),
        ("team rocket", "Team Rocket"),
        ("league", "Pokémon League")
    ]
    
    formatted_sets = {}
    tr_id = 1
    
    # Store tr_ids in order to build the navigation sequence
    tr_ids_list = []
    
    for key, display_name in sheets:
        actual_sheetname = sheet_mapping.get(key)
        if not actual_sheetname:
            print(f"Warning: sheet for '{key}' not found in Excel file.")
            continue
            
        sheet = wb[actual_sheetname]
        rows = list(sheet.iter_rows(values_only=True))
        
        # Iterating through rows in blocks of 19
        for r_idx in range(0, len(rows), 19):
            if r_idx >= len(rows):
                break
                
            trainer_row = rows[r_idx]
            if not trainer_row or len(trainer_row) == 0:
                continue
            trainer_name = trainer_row[0]
            if not trainer_name:
                continue
                
            trainer_name = str(trainer_name).strip()
            
            # Check if there is at least a level row (row 9, index 8)
            if r_idx + 8 >= len(rows):
                break
                
            species_row = rows[r_idx + 1]
            level_row = rows[r_idx + 8]
            type_row = rows[r_idx + 9] if r_idx + 9 < len(rows) else None
            ability_row = rows[r_idx + 10] if r_idx + 10 < len(rows) else None
            item_row = rows[r_idx + 11] if r_idx + 11 < len(rows) else None
            nature_row = rows[r_idx + 12] if r_idx + 12 < len(rows) else None
            
            move_rows = []
            for m_offset in range(13, 17):
                if r_idx + m_offset < len(rows):
                    move_rows.append(rows[r_idx + m_offset])
                    
            # Parse the Pokémon (up to 6 columns, B through G, indices 1 to 6)
            parsed_any = False
            species_counts = {} # Keep track of species counts within this team
            
            for col_idx in range(1, 7):
                if col_idx >= len(species_row):
                    break
                    
                raw_species = species_row[col_idx]
                if not raw_species:
                    continue
                    
                species = clean_species(raw_species)
                if not species:
                    continue
                    
                # Parse details
                raw_level = level_row[col_idx] if col_idx < len(level_row) else 100
                try:
                    level = int(float(raw_level))
                except:
                    level = 100
                    
                ability = clean_ability(ability_row[col_idx]) if ability_row and col_idx < len(ability_row) else "None"
                item = clean_item(item_row[col_idx]) if item_row and col_idx < len(item_row) else "None"
                nature = clean_nature(nature_row[col_idx]) if nature_row and col_idx < len(nature_row) else "Serious"
                
                moves = []
                for m_row in move_rows:
                    if col_idx < len(m_row):
                        m_name = clean_move(m_row[col_idx])
                        if m_name:
                            moves.append(m_name)
                            
                sub_index = col_idx - 1
                
                # Deduplicate key name for multiple Pokémon of the same species
                count = species_counts.get(species, 0)
                species_counts[species] = count + 1
                asterisks = "*" * count
                set_name = f"Lvl {level}{asterisks} {trainer_name} "
                
                # Add to formatted_sets
                if species not in formatted_sets:
                    formatted_sets[species] = {}
                    
                formatted_sets[species][set_name] = {
                    "level": level,
                    "tr_id": tr_id,
                    "ai": 7,
                    "battle_type": "Singles",
                    "reward_item": "",
                    "form": 0,
                    "item": item,
                    "ivs": {
                        "hp": 31,
                        "at": 31,
                        "df": 31,
                        "sa": 31,
                        "sd": 31,
                        "sp": 31
                    },
                    "evs": {
                        "hp": 0,
                        "at": 0,
                        "df": 0,
                        "sa": 0,
                        "sd": 0,
                        "sp": 0
                    },
                    "nature": nature,
                    "moves": moves,
                    "sub_index": sub_index,
                    "ability": ability,
                    "gender": "Male",
                    "location": display_name,
                    "spriteId": None,
                    "orientation": None
                }
                parsed_any = True
                
            if parsed_any:
                tr_ids_list.append(tr_id)
                tr_id += 1
                
    # Build order dictionary
    order = {}
    for i in range(len(tr_ids_list)):
        curr_id = tr_ids_list[i]
        prev_id = tr_ids_list[i - 1] if i > 0 else 0
        next_id = tr_ids_list[i + 1] if i < len(tr_ids_list) - 1 else 0
        order[curr_id] = {
            "next": next_id,
            "prev": prev_id
        }
        
    # Write backups/fireredimproved.js
    output_path = "backups/fireredimproved.js"
    backup_data = {
        "title": "Rosso Fuoco Migliorato By Musoki",
        "formatted_sets": formatted_sets,
        "poks": {},
        "moves": {},
        "order": order
    }
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("backup_data = ")
        json.dump(backup_data, f, indent=4)
        f.write(";\n")
        
    print(f"Successfully generated {output_path} with {len(tr_ids_list)} trainer teams.")

if __name__ == "__main__":
    main()
