import openpyxl
import re
import json
import zipfile
import xml.etree.ElementTree as ET
import os

SPECIES_CORRECTIONS = {
    "Skamory": "Skarmory",
    "Snubbul": "Snubbull",
    "Farfetchd": "Farfetch'd",
    "Sirfetchd": "Sirfetch'd",
    "Mr-Mime": "Mr. Mime",
    "Mime-Jr": "Mime Jr.",
    "Mr-Rime": "Mr. Rime",
    "Magnazone": "Magnezone",
    "Mismagius (Gengar)": "Mismagius",
}

MOVE_CORRECTIONS = {
    "ExtremeSpeed": "Extreme Speed",
    "Poisonpowder": "Poison Powder",
    "Sonicboom": "Sonic Boom",
    "ThunderPunch": "Thunder Punch",
    "Thundershock": "Thunder Shock",
    "U-Turn": "U-turn",
    "Dynamicpunch": "Dynamic Punch",
}

ITEM_CORRECTIONS = {
    "Never-Melt Ice": "NeverMelt Ice",
    "Never-Melt-Ice": "NeverMelt Ice",
}

def clean_species(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r"\s*[\(\[].*?[\)\]]", "", name).strip()
    return SPECIES_CORRECTIONS.get(name, name)

def clean_move(name):
    if not name:
        return ""
    name = str(name).strip()
    if name == "-" or name.lower() == "none" or name.lower() == "null":
        return ""
    return MOVE_CORRECTIONS.get(name, name)

def clean_item(name):
    if not name:
        return "None"
    name = str(name).strip()
    if name.lower() == "none" or name.lower() == "null" or name == "-":
        return "None"
    return ITEM_CORRECTIONS.get(name, name)

def clean_nature(name):
    if not name:
        return "Serious"
    name = str(name).strip()
    return name

def clean_ability(name):
    if not name:
        return "None"
    return str(name).strip()

def normalize_str(s):
    return re.sub(r'[^a-zA-Z0-9]', '', str(s)).lower()

def get_docx_paragraphs(path):
    with zipfile.ZipFile(path) as docx:
        tree = ET.fromstring(docx.read('word/document.xml'))
        namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paras = []
        for p in tree.iterfind('.//w:p', namespaces):
            text = ''.join([node.text for node in p.iter() if node.text])
            if text.strip():
                paras.append(text.strip())
        return paras

def parse_docx_pokemon_sets(path):
    if not os.path.exists(path):
        return []
    paras = get_docx_paragraphs(path)
    sets = []
    curr = []
    for p in paras:
        if p.startswith('Calcolatore'): continue
        is_header = False
        if re.search(r'\([A-Za-z0-9\.\-\s\']+\)', p) or re.search(r'Leader|Rival|Executive|Elite|Champion|Trainer|Elder|Eusine|Kiyo|Kimono|Ace|Psychic|Fisherman|Grunt', p):
            if not any(p.startswith(x) for x in ['- ', 'IVs', 'EVs', 'Ability', 'Level', 'Shiny', 'Gigantamax', 'Tera Type', 'Type:', 'Nature']):
                if ' - ' in p or ' @ ' in p or '(' in p or any(title in p for title in ['Leader', 'Rival', 'Executive', 'Elite', 'Champion', 'Trainer', 'Elder', 'Eusine', 'Kiyo', 'Kimono', 'Ace', 'Psychic', 'Fisherman']):
                    is_header = True

        if is_header:
            if curr:
                sets.append(curr)
                curr = []
        if curr or is_header:
            curr.append(p)
    if curr:
        sets.append(curr)

    parsed_sets = []
    for s in sets:
        header = s[0]
        item = "None"
        if " @ " in header:
            header_part, item = header.split(" @ ", 1)
            item = item.strip()
        else:
            header_part = header

        species_match = re.search(r'\(([A-Za-z0-9\.\-\s\']+)\)', header_part)
        raw_species = ""
        if species_match:
            raw_species = species_match.group(1).strip()
        else:
            parts = header_part.split(' - ')
            if len(parts) > 1:
                raw_species = parts[1].strip()

        trainer_raw = header_part.split(' - ')[0].strip() if ' - ' in header_part else ""

        ability = "None"
        level = 100
        nature = "Hardy"
        ivs = {"hp": 31, "at": 31, "df": 31, "sa": 31, "sd": 31, "sp": 31}
        evs = {"hp": 0, "at": 0, "df": 0, "sa": 0, "sd": 0, "sp": 0}
        moves = []

        for line in s[1:]:
            line = line.strip()
            if line.startswith('Ability:'):
                ability = line.replace('Ability:', '').strip()
            elif line.startswith('Level:'):
                try:
                    level = int(line.replace('Level:', '').strip())
                except:
                    pass
            elif 'Nature' in line:
                nature = line.replace('Nature', '').strip()
            elif line.startswith('IVs:'):
                iv_str = line.replace('IVs:', '').strip()
                parts = iv_str.split('/')
                stat_map = {'hp': 'hp', 'atk': 'at', 'def': 'df', 'fdef': 'df', 'spa': 'sa', 'spd': 'sd', 'spe': 'sp'}
                for pt in parts:
                    pt = pt.strip()
                    m = re.match(r'(\d+)\s*([A-Za-z]+)', pt)
                    if m:
                        val = int(m.group(1))
                        stat = m.group(2).lower()
                        if stat in stat_map:
                            ivs[stat_map[stat]] = val
            elif line.startswith('EVs:'):
                ev_str = line.replace('EVs:', '').strip()
                parts = ev_str.split('/')
                stat_map = {'hp': 'hp', 'atk': 'at', 'def': 'df', 'fdef': 'df', 'spa': 'sa', 'spd': 'sd', 'spe': 'sp'}
                for pt in parts:
                    pt = pt.strip()
                    m = re.match(r'(\d+)\s*([A-Za-z]+)', pt)
                    if m:
                        val = int(m.group(1))
                        stat = m.group(2).lower()
                        if stat in stat_map:
                            evs[stat_map[stat]] = val
            elif line.startswith('- '):
                move = line[2:].strip()
                if move:
                    moves.append(move)

        parsed_sets.append({
            "header": header,
            "trainer": trainer_raw,
            "species": clean_species(raw_species),
            "item": clean_item(item),
            "ability": clean_ability(ability),
            "level": level,
            "nature": clean_nature(nature),
            "ivs": ivs,
            "evs": evs,
            "moves": [clean_move(m) for m in moves]
        })

    return parsed_sets

def main():
    excel_candidates = [
        "Docs/2.0/Documento Lotte Pokemon Heart Gold MA MIGLIORATO 2.0.xlsx",
        "Documento Lotte Pokemon Heart Gold MA MIGLIORATO 2.0.xlsx",
        "Documento Lotte Pokemon Heart Gold MA MIGLIORATO.xlsx"
    ]
    excel_path = next((p for p in excel_candidates if os.path.exists(p)), None)
    if not excel_path:
        raise FileNotFoundError("Could not find Heart Gold trainer Excel file.")
    print(f"Loading battles Excel from {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, read_only=True)

    docx_candidates = [
        "Docs/2.0/Squadre da portare sul Calcolatore di Pokémon Showdown.docx",
        "Docs/Squadre da portare sul Calcolatore di Pokémon Showdown.docx"
    ]
    docx_path = next((p for p in docx_candidates if os.path.exists(p)), None)
    docx_sets = []
    docx_lookup = {}
    if docx_path:
        print(f"Loading Showdown teams & IVs from {docx_path}...")
        docx_sets = parse_docx_pokemon_sets(docx_path)
        for d in docx_sets:
            tr = d['trainer']
            sp = d['species']
            lvl = d['level']
            key = (normalize_str(tr), normalize_str(sp), lvl)
            if key not in docx_lookup:
                docx_lookup[key] = []
            docx_lookup[key].append(d)
        print(f"Loaded {len(docx_sets)} Pokémon sets with IVs from Showdown document.")

    sheet_mapping = {
        "johto": None,
        "rival": None,
        "rocket": None,
        "mini boss": None,
        "route 27": None,
        "league": None,
        "kanto": None,
        "red": None
    }
    
    for actual_name in wb.sheetnames:
        clean_name = actual_name.lower().replace(" ", "").replace("é", "e")
        for key in sheet_mapping.keys():
            if key.replace(" ", "") in clean_name:
                sheet_mapping[key] = actual_name
                break
                
    sheets = [
        ("johto", "Johto Gym Leader"),
        ("rival", "Rival"),
        ("rocket", "Team Rocket"),
        ("mini boss", "Mini Boss Fight Obligatory"),
        ("route 27", "Route 27 / Route 26 Fight"),
        ("league", "Pokémon League"),
        ("kanto", "Kanto Gym Leader"),
        ("red", "Red Mt. Silver")
    ]
    
    formatted_sets = {}
    tr_id = 1
    tr_ids_list = []
    
    for key, display_name in sheets:
        actual_sheetname = sheet_mapping.get(key)
        if not actual_sheetname:
            print(f"Warning: sheet for '{key}' not found in Excel file.")
            continue
            
        sheet = wb[actual_sheetname]
        rows = list(sheet.iter_rows(values_only=True))

        for r_idx in range(0, len(rows), 19):
            if r_idx >= len(rows):
                break
                
            trainer_row = rows[r_idx]
            trainer_name = trainer_row[0]
            if not trainer_name:
                continue
                
            trainer_name = str(trainer_name).strip()
            
            if r_idx + 8 >= len(rows):
                break
                
            species_row = rows[r_idx + 1]
            level_row = rows[r_idx + 8]
            ability_row = rows[r_idx + 10] if r_idx + 10 < len(rows) else None
            item_row = rows[r_idx + 11] if r_idx + 11 < len(rows) else None
            nature_row = rows[r_idx + 12] if r_idx + 12 < len(rows) else None
            
            move_rows = []
            for m_offset in range(13, 17):
                if r_idx + m_offset < len(rows):
                    move_rows.append(rows[r_idx + m_offset])
                    
            parsed_any = False
            species_counts = {}
            
            for col_idx in range(1, 7):
                if col_idx >= len(species_row):
                    break
                    
                raw_species = species_row[col_idx]
                if not raw_species:
                    continue
                    
                species = clean_species(raw_species)
                if not species:
                    continue
                    
                raw_level = level_row[col_idx] if col_idx < len(level_row) else 100
                try:
                    level = int(float(raw_level))
                except:
                    level = 100
                    
                ability = clean_ability(ability_row[col_idx]) if ability_row and col_idx < len(ability_row) else "None"
                item = clean_item(item_row[col_idx]) if item_row and col_idx < len(item_row) else "None"
                nature = clean_nature(nature_row[col_idx]) if nature_row and col_idx < len(nature_row) else "Serious"
                
                moves = []
                for m_row in move_rows:
                    if col_idx < len(m_row):
                        m_name = clean_move(m_row[col_idx])
                        if m_name:
                            moves.append(m_name)
                            
                # Match IVs & EVs from docx if available
                ivs = {"hp": 31, "at": 31, "df": 31, "sa": 31, "sd": 31, "sp": 31}
                evs = {"hp": 0, "at": 0, "df": 0, "sa": 0, "sd": 0, "sp": 0}
                
                if docx_lookup:
                    clean_tr = normalize_str(trainer_name)
                    clean_sp = normalize_str(species)
                    sp_candidates = [clean_sp]
                    if clean_sp == 'skamory': sp_candidates.append('skarmory')
                    if clean_sp == 'snubbul': sp_candidates.append('snubbull')
                    if clean_sp == 'mismagius': sp_candidates.append('gengar')
                    if '(a)' in str(raw_species).lower(): sp_candidates.append(clean_sp.replace('a', ''))
                    if '(g)' in str(raw_species).lower(): sp_candidates.append(clean_sp.replace('g', ''))

                    found_dx = None
                    for cand_sp in sp_candidates:
                        for (k_tr, k_sp, k_lvl), d_list in docx_lookup.items():
                            if k_sp == cand_sp and (k_lvl == level or abs(k_lvl - level) <= 1):
                                if (clean_tr in k_tr or k_tr in clean_tr or 
                                    any(part in k_tr for part in ['falkner', 'bugsy', 'whitney', 'morty', 'chuck', 'jasmine', 'pryce', 'clair',
                                                                 'will', 'koga', 'bruno', 'karen', 'lance',
                                                                 'proton', 'petrel', 'ariana', 'archer',
                                                                 'brock', 'misty', 'surge', 'erika', 'janine', 'sabrina', 'blaine', 'blue',
                                                                 'red', 'li', 'eusine', 'kiyo', 'zuki', 'naoko', 'miki', 'sajo', 'kuni',
                                                                 'megan', 'blake', 'brian', 'eli', 'reena', 'scott', 'vernon', 'joyce', 'gaven', 'jake', 'jamie'] if part in clean_tr)):
                                    if d_list:
                                        found_dx = d_list[0]
                                        break
                        if found_dx:
                            break
                            
                    if found_dx:
                        ivs = found_dx["ivs"]
                        evs = found_dx["evs"]
                        if item == "None" and found_dx["item"] != "None":
                            item = found_dx["item"]
                        if ability == "None" and found_dx["ability"] != "None":
                            ability = found_dx["ability"]
                        if len(moves) == 0 and len(found_dx["moves"]) > 0:
                            moves = found_dx["moves"]

                sub_index = col_idx - 1
                
                # Deduplicate key name for multiple Pokémon of the same species
                count = species_counts.get(species, 0)
                species_counts[species] = count + 1
                asterisks = "*" * count
                set_name = f"Lvl {level}{asterisks} {trainer_name} "
                
                # Add to formatted_sets
                if species not in formatted_sets:
                    formatted_sets[species] = {}
                    
                formatted_sets[species][set_name] = {
                    "level": level,
                    "tr_id": tr_id,
                    "ai": 7,
                    "battle_type": "Singles",
                    "reward_item": "",
                    "form": 0,
                    "item": item,
                    "ivs": ivs,
                    "evs": evs,
                    "nature": nature,
                    "moves": moves,
                    "sub_index": sub_index,
                    "ability": ability,
                    "gender": "Male",
                    "location": display_name,
                    "spriteId": None,
                    "orientation": None
                }
                parsed_any = True
                
            if parsed_any:
                tr_ids_list.append(tr_id)
                tr_id += 1
                
    # Build order dictionary
    order = {}
    for i in range(len(tr_ids_list)):
        curr_id = tr_ids_list[i]
        prev_id = tr_ids_list[i - 1] if i > 0 else 0
        next_id = tr_ids_list[i + 1] if i < len(tr_ids_list) - 1 else 0
        order[curr_id] = {
            "next": next_id,
            "prev": prev_id
        }
        
    # Write backups/hgimproved.js
    output_path = "backups/hgimproved.js"
    backup_data = {
        "title": "Heart Gold Migliorato 2.0",
        "formatted_sets": formatted_sets,
        "poks": {},
        "moves": {},
        "order": order
    }
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("backup_data = ")
        json.dump(backup_data, f, indent=4)
        f.write(";\n")
        
    print(f"Successfully generated {output_path} with {len(tr_ids_list)} trainer teams and {sum(len(v) for v in formatted_sets.values())} Pokemon sets.")

if __name__ == "__main__":
    main()
