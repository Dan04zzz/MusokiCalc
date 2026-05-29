import openpyxl
import re
import json
import datetime

SPECIES_CORRECTIONS = {
    "Skamory": "Skarmory",
    "Snubbul": "Snubbull",
    "Farfetchd": "Farfetch'd",
    "Sirfetchd": "Sirfetch'd",
    "Mr-Mime": "Mr. Mime",
    "Mime-Jr": "Mime Jr.",
    "Mr-Rime": "Mr. Rime",
}

def clean_species(name):
    if not name:
        return ""
    name = str(name).strip()
    
    # Translate some common spelling issues
    # Strip any brackets or parentheses (e.g. eggs, items, etc.)
    name_clean = re.sub(r"\s*[\(\[].*?[\)\]]", "", name).strip()
    
    # Spelling fixes
    if name_clean in SPECIES_CORRECTIONS:
        return SPECIES_CORRECTIONS[name_clean]
        
    lower_name = name_clean.lower().replace(" ", "").replace("-", "").replace(".", "").replace("'", "")
    
    if lower_name == "nidoranf" or lower_name == "nidoranfemmina" or lower_name == "nidoranfemminile":
        return "Nidoran-F"
    if lower_name == "nidoranm" or lower_name == "nidoranmaschio" or lower_name == "nidoranmaschile":
        return "Nidoran-M"
    if lower_name == "hooh":
        return "Ho-Oh"
    if lower_name == "porygonz":
        return "Porygon-Z"
    if lower_name == "porygon2":
        return "Porygon2"
    if lower_name == "mrmime":
        return "Mr. Mime"
    if lower_name == "mimejr":
        return "Mime Jr."
    if lower_name == "farfetchd":
        return "Farfetch'd"
    if lower_name == "giratinaorigin":
        return "Giratina-Origin"
    if lower_name == "rotomwash":
        return "Rotom-Wash"
    if lower_name == "rotomheat":
        return "Rotom-Heat"
    if lower_name == "rotomfrost":
        return "Rotom-Frost"
    if lower_name == "rotomfan":
        return "Rotom-Fan"
    if lower_name == "rotommow":
        return "Rotom-Mow"
    if lower_name == "deoxysattack":
        return "Deoxys-Attack"
    if lower_name == "deoxysdefense":
        return "Deoxys-Defense"
    if lower_name == "deoxysspeed":
        return "Deoxys-Speed"
    if lower_name == "wormadamtrash":
        return "Wormadam-Trash"
    if lower_name == "wormadamsandy":
        return "Wormadam-Sandy"
    if lower_name == "shayminsky":
        return "Shaymin-Sky"
    
    return name_clean

def clean_level(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        # Excel date translation for DD/MM formats (e.g. 2-4 interpreted as 2nd of April, day=2, month=4)
        return f"{val.day}-{val.month}"
    
    val_str = str(val).strip()
    if not val_str:
        return ""
        
    # Check if it matches a date-like string format
    date_match = re.match(r"(\d{4})-(\d{2})-(\d{2})", val_str)
    if date_match:
        day = int(date_match.group(3))
        month = int(date_match.group(2))
        return f"{day}-{month}"
        
    try:
        f = float(val_str)
        if f.is_integer():
            return str(int(f))
        return str(f)
    except ValueError:
        pass
        
    return val_str

def clean_rate(val):
    if val is None:
        return ""
    try:
        f = float(val)
        if f <= 1.0:
            return f"{int(round(f * 100))}%"
        else:
            return f"{int(round(f))}%"
    except ValueError:
        pass
    val_str = str(val).strip()
    if val_str and not val_str.endswith('%'):
        try:
            float(val_str)
            return val_str + "%"
        except ValueError:
            pass
    return val_str

def main():
    wb = openpyxl.load_workbook("Pokémon Rosso Fuoco MA MIGLIORATO Encounters.xlsx", read_only=True)
    
    # 1. Parse Wild Pokémon
    wild_sheet = wb['Wild Pok\xe9mon' if 'Wild Pok\xe9mon' in wb.sheetnames else wb.sheetnames[0]]
    wild_rows = list(wild_sheet.iter_rows(values_only=True))
    
    # Define row blocks for FireRed
    encounter_types = [
        ("Grass", 3, 15),
        ("Surfing", 15, 20),
        ("Old Rod", 20, 22),
        ("Good Rod", 22, 25),
        ("Super Rod", 25, 30),
        ("Rock Smash", 30, 35)
    ]
    
    routes_data = []
    num_cols = len(wild_rows[0])
    
    for col_idx in range(2, num_cols, 3):
        route_name = wild_rows[0][col_idx]
        if not route_name:
            continue
        route_name = str(route_name).strip()
        
        route_encounters = {}
        for label, start_r, end_r in encounter_types:
            list_pokes = []
            for r in range(start_r, end_r):
                if r >= len(wild_rows):
                    break
                poke = wild_rows[r][col_idx]
                lvl = wild_rows[r][col_idx + 1]
                rate = wild_rows[r][col_idx + 2]
                
                cleaned_poke = clean_species(poke)
                if cleaned_poke:
                    list_pokes.append({
                        "species": cleaned_poke,
                        "level": clean_level(lvl),
                        "rate": clean_rate(rate)
                    })
            if list_pokes:
                clean_label = label
                if "Old Rod" in clean_label or "Old Road" in clean_label:
                    clean_label = "Old Rod"
                elif "Good Rod" in clean_label or "Good Road" in clean_label:
                    clean_label = "Good Rod"
                elif "Super Rod" in clean_label or "Super Road" in clean_label:
                    clean_label = "Super Rod"
                route_encounters[clean_label] = list_pokes
                
        if route_encounters:
            routes_data.append({
                "name": route_name,
                "encounters": route_encounters
            })
            
    # 2. Parse Gift, Static, Trade, Game Corner Sheet
    gift_sheet = wb['Gift, Static, Trade, Game Corne']
    gift_rows = list(gift_sheet.iter_rows(values_only=True))
    
    gifts = []
    statics = []
    trades = []
    game_corner = []
    
    for r in range(3, len(gift_rows)):
        row = gift_rows[r]
        if r >= len(gift_rows):
            break
            
        # Gift: col 0, 1, 2
        if len(row) > 0 and row[0]:
            gifts.append({
                "species": clean_species(row[0]),
                "location": str(row[1]).strip() if row[1] else "",
                "level": clean_level(row[2])
            })
            
        # Static: col 3, 4, 5
        if len(row) > 3 and row[3]:
            statics.append({
                "species": clean_species(row[3]),
                "location": str(row[4]).strip() if row[4] else "",
                "level": clean_level(row[5])
            })
            
        # Trade: col 6, 7, 8
        if len(row) > 7 and row[7]:
            trades.append({
                "give": clean_species(row[6]),
                "receive": clean_species(row[7]),
                "location": str(row[8]).strip() if row[8] else ""
            })
            
        # Game Corner Col 9-11
        if len(row) > 9 and row[9]:
            game_corner.append({
                "species": clean_species(row[9]),
                "price": str(row[10]).strip() if row[10] else "",
                "level": clean_level(row[11])
            })
            
    # Combine everything
    encounters_database = {
        "routes": routes_data,
        "gifts": gifts,
        "statics": statics,
        "trades": trades,
        "game_corner": game_corner
    }
    
    # Write to backups/fireredimproved_encounters.js
    output_path = "backups/fireredimproved_encounters.js"
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("window.fireredimproved_encounters = ")
        json.dump(encounters_database, f, indent=4)
        f.write(";\n")
        
    print(f"Successfully generated {output_path}")
    print(f"Parsed {len(routes_data)} routes.")
    print(f"Parsed {len(gifts)} gifts, {len(statics)} statics, {len(trades)} trades, {len(game_corner)} game corner mons.")

if __name__ == "__main__":
    main()
